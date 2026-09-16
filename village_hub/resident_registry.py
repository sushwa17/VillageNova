"""Excel-backed resident registry and per-resident delivery history."""
import json
import shutil
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook

DATA_DIR = Path(__file__).parent / "data"
RESIDENTS_FILE = DATA_DIR / "residents.xlsx"
LEGACY_FILE = DATA_DIR / "residents.json"
RESIDENT_HEADERS = [
    "ID", "Name", "Location", "Phone", "Channel", "Language",
    "Categories", "LastSent",
]
DELIVERY_HEADERS = [
    "Timestamp", "ResidentID", "Name", "Location", "Phone", "Channel",
    "Categories", "Message", "Status",
]

CATEGORIES = {
    "farmer": "Farmers",
    "cattle": "Cattle owners",
    "irrigation": "Irrigation and water",
    "livelihoods": "Workers and small businesses",
    "women": "Women support",
    "student": "Students and education",
    "health": "Health information",
    "public_services": "Panchayat and public services",
    "general": "General village updates",
}


def _create_workbook() -> None:
    workbook = Workbook()
    residents = workbook.active
    residents.title = "Residents"
    residents.append(RESIDENT_HEADERS)
    delivery_log = workbook.create_sheet("DeliveryLog")
    delivery_log.append(DELIVERY_HEADERS)
    workbook.save(RESIDENTS_FILE)


def _ensure_workbook() -> None:
    DATA_DIR.mkdir(exist_ok=True)
    if not RESIDENTS_FILE.exists():
        _create_workbook()
        _migrate_legacy_json()
        return
    try:
        workbook = load_workbook(RESIDENTS_FILE, read_only=True)
        workbook.close()
    except (BadZipFile, OSError):
        backup = RESIDENTS_FILE.with_name(
            f"{RESIDENTS_FILE.stem}.corrupt-{datetime.now():%Y%m%d-%H%M%S}{RESIDENTS_FILE.suffix}"
        )
        shutil.move(RESIDENTS_FILE, backup)
        print(f"[resident_registry] Invalid workbook moved to {backup.name}; rebuilding it.")
        _create_workbook()
        _migrate_legacy_json()


def _migrate_legacy_json() -> None:
    if not LEGACY_FILE.exists():
        return
    try:
        legacy_residents = json.loads(LEGACY_FILE.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return
    if not legacy_residents:
        return
    workbook = load_workbook(RESIDENTS_FILE)
    sheet = workbook["Residents"]
    for resident in legacy_residents:
        name = resident.get("name") or resident.get("nae") or "Unknown resident"
        categories = [
            category if category in CATEGORIES else "general"
            for category in resident.get("categories", [])
        ]
        sheet.append([
            resident["id"], name, resident["location"],
            resident["phone"], resident["channel"], resident["language"],
            ",".join(sorted(set(categories))), resident.get("last_sent"),
        ])
    workbook.save(RESIDENTS_FILE)


def add_resident(record: dict) -> int:
    _ensure_workbook()
    workbook = load_workbook(RESIDENTS_FILE)
    sheet = workbook["Residents"]
    resident_id = max((row[0].value or 0 for row in sheet.iter_rows(min_row=2)), default=0) + 1
    sheet.append([
        resident_id, record["name"], record["location"], record["phone"],
        record["channel"], record["language"], ",".join(record["categories"]), None,
    ])
    workbook.save(RESIDENTS_FILE)
    return resident_id


def get_residents() -> list[dict]:
    _ensure_workbook()
    workbook = load_workbook(RESIDENTS_FILE, read_only=True)
    sheet = workbook["Residents"]
    headers = [cell.value for cell in sheet[1]]
    residents = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        resident = dict(zip(headers, row))
        categories = [item.strip() for item in (resident.get("Categories") or "").split(",") if item.strip()]
        residents.append({
            "id": resident["ID"], "name": resident["Name"], "location": resident["Location"],
            "phone": resident["Phone"], "channel": resident["Channel"],
            "language": resident["Language"], "categories": categories,
            "last_sent": resident["LastSent"],
        })
    return residents


def mark_sent(resident: dict, message: str, status: str) -> None:
    """Update the resident row and append an individual delivery record."""
    _ensure_workbook()
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    workbook = load_workbook(RESIDENTS_FILE)
    residents = workbook["Residents"]
    for row in residents.iter_rows(min_row=2):
        if row[0].value == resident["id"]:
            row[7].value = timestamp
            break
    workbook["DeliveryLog"].append([
        timestamp, resident["id"], resident["name"], resident["location"],
        resident["phone"], resident["channel"], ",".join(resident["categories"]),
        message, status,
    ])
    workbook.save(RESIDENTS_FILE)
