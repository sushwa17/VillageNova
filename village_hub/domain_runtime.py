"""Runtime shared by the independently runnable village domain agents."""
import json
import os
import shutil
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
import requests
from dotenv import load_dotenv

load_dotenv(Path(__file__).parents[1] / "krishi_agent" / ".env")

DOMAIN_DEFINITIONS = {
    "cattle": {
        "title": "Cattle Agent",
        "fields": [
            ("cattle_type", "Which type of cattle (cow/buffalo/goat/sheep): "),
            ("animals", "Animals (cow/buffalo/goat/sheep): "),
            ("count", "Approximate number of animals: "),
            ("concern", "Main concern (fodder/health/milk/vaccination): "),
        ],
        "messages": [
            "Daily cattle tip: for {cattle_type}, keep clean water and shade ready and watch {concern} closely.",
            "Feeding note: use green fodder, dry roughage, and mineral support; clean water comes first.",
            "Mark the next vaccination or deworming date with your local veterinary worker.",
            "If an animal stops eating, has high fever, or cannot stand, call a vet immediately.",
        ],
    },
    "irrigation": {
        "title": "Irrigation Agent",
        "fields": [
            ("source", "Water source (borewell/canal/tank/river): "),
            ("crop", "Main crop or area served: "),
            ("irrigation_method", "Irrigation method (drip/sprinkler/flood/furrow): "),
            ("equipment", "Equipment (pump/drip/sprinkler/none): "),
            ("stage", "Crop stage (sowing/seedling/flowering/fruiting/maturity): "),
        ],
        "messages": [
            "Water update: for {crop} from {source}, check the local turn before starting the pump.",
            "Use {irrigation_method} scheduling to match the crop stage and cut water waste.",
            "Before running {equipment}, check cables, pipes, earthing, and leaks.",
            "Irrigate early morning or evening when possible to reduce evaporation and save energy.",
        ],
    },
    "livelihoods": {
        "title": "Livelihoods Agent",
        "fields": [
            ("work_type", "Work or skill (driver/tailor/shop/artisan/labour/other): "),
            ("work", "Work or skill (driver/tailor/shop/artisan/labour/other): "),
            ("need", "Information needed (jobs/training/loans/market): "),
        ],
        "messages": [
            "Livelihood note: for {work_type}, check the local panchayat and skill-centre boards for trusted work.",
            "For {need}, compare nearby demand, wages, and verified options before paying any fee.",
            "Keep a small record of work, payments, and customers to protect your income.",
        ],
    },
    "women": {
        "title": "Women Support Agent",
        "fields": [
            ("support_type", "Support needed (health/SHG/training/safety/childcare): "),
            ("interest", "Main interest (health/SHG/training/safety/childcare): "),
            ("group", "Self-help group or community group (optional): "),
        ],
        "messages": [
            "Support alert: for {support_type}, ask the ASHA, ANM, Anganwadi, or panchayat worker about the next local service date.",
            "For self-help groups like {group}, keep records of savings, meetings, and loans in a safe place.",
            "If there is immediate danger or violence, move to safety and contact a trusted person or local emergency support.",
        ],
    },
    "health": {
        "title": "Health Agent",
        "fields": [("topic", "Health topic (nutrition/maternal/child/elderly/sanitation): "), ("age_group", "Person or age group needing information: ")],
        "messages": [
            "Health update for {topic}, especially for {age_group}: confirm dates with the government health centre or ASHA worker.",
            "Do not delay urgent care for severe breathing trouble, chest pain, heavy bleeding, unconsciousness, or serious injury.",
            "Keep prescriptions and vaccination records together and do not share private medical details publicly.",
        ],
    },
    "education": {
        "title": "Education Agent",
        "fields": [
            ("school_name", "School or college name: "),
            ("learner", "Learner (child/adult/student): "),
            ("ai_focus", "AI topic or school innovation focus (chatbot/robotics/writing/math/project): "),
            ("need", "Need (school/scholarship/exam/skills/digital): "),
        ],
        "messages": [
            "School update: {learner} at {school_name} should check the notice board or official portal for {need} updates.",
            "AI school note: explore {ai_focus} in a practical way to build creativity, problem-solving, and digital skills.",
            "Ask the teacher or education office about scholarships and deadlines; never pay an unverified agent.",
            "Set a daily study time and keep school records and application numbers safe.",
        ],
    },
    "public_services": {
        "title": "Public Services Agent",
        "fields": [("service", "Service needed (ration/pension/water/power/documents/roads): "), ("ward", "Ward, hamlet, or local area: ")],
        "messages": [
            "Public-service update for {service} in {ward}: verify the latest notice with the panchayat office.",
            "Keep application receipts and acknowledgement numbers; do not pay unofficial fees for government services.",
            "Report unsafe roads, outages, or water problems through the official local complaint channel.",
        ],
    },
    "village_updates": {
        "title": "Village Updates Agent",
        "fields": [("interests", "Interests (news/health/weather/school/schemes): ")],
        "messages": [
            "Village update for {interests}: use verified panchayat, school, health-centre, and department notices.",
            "Check the date and source before forwarding a message; do not forward rumours or unverified emergency claims.",
            "Keep important local helpline numbers saved and share urgent verified notices responsibly.",
        ],
    },
}

FARMER_UPDATE = (
    "Farmer update: check today's weather and crop conditions before field work. "
    "Use the Krishi Setu advisory for crop-specific guidance."
)


LOCALIZED_MESSAGES = {
    "cattle": {
        "hi": [
            "{animals} के लिए साफ पानी और छाया रखें; {concern} पर ध्यान दें।",
            "अगले टीकाकरण या कृमिनाशक की तारीख स्थानीय पशु स्वास्थ्यकर्मी के साथ लिख लें।",
            "अगर पशु खाना छोड़ दे, तेज बुखार हो या खड़ा न हो पाए, तो तुरंत पशु चिकित्सक से संपर्क करें।",
        ],
        "te": [
            "{animals}కు శుభ్రమైన నీరు, నీడ అందించండి; {concern}పై శ్రద్ధ వహించండి.",
            "తదుపరి టీకా లేదా పురుగుల మందు తేదీని స్థానిక పశువైద్య సిబ్బందితో నమోదు చేసుకోండి.",
            "జంతువు తినకపోతే, ఎక్కువ జ్వరం ఉంటే లేదా నిలబడలేకపోతే వెంటనే పశువైద్యుడిని సంప్రదించండి.",
        ],
    },
    "irrigation": {
        "hi": [
            "{source} से {crop} की सिंचाई के लिए पंप चलाने से पहले स्थानीय पानी की बारी जांचें।",
            "{equipment} चलाने से पहले तार, अर्थिंग, पाइप और रिसाव जांचें।",
            "संभव हो तो सुबह जल्दी या शाम को पानी दें ताकि पानी की बचत हो।",
        ],
        "te": [
            "{source} ద్వారా {crop}కు నీరు పెట్టే ముందు స్థానిక నీటి వంతును తనిఖీ చేయండి.",
            "{equipment} ఉపయోగించే ముందు కేబుళ్లు, ఎర్తింగ్, పైపులు, లీకేజీలను తనిఖీ చేయండి.",
            "సాధ్యమైనప్పుడు ఉదయం లేదా సాయంత్రం నీరు పెట్టి నీటి ఆవిరీభవనాన్ని తగ్గించండి.",
        ],
    },
    "livelihoods": {
        "hi": [
            "{work} के लिए आजीविका सूचना: प्रमाणित अवसरों के लिए पंचायत और कौशल केंद्र के सूचना बोर्ड देखें।",
            "{need} के लिए दस्तावेज देने या शुल्क भरने से पहले सरकारी बैंक या पंजीकृत केंद्र से पुष्टि करें।",
            "अपनी आय की सुरक्षा के लिए काम, भुगतान और ग्राहक के ऑर्डर का सरल रिकॉर्ड रखें।",
        ],
        "te": [
            "{work}కు జీవనోపాధి సమాచారం: ధృవీకరించిన అవకాశాల కోసం పంచాయతీ, నైపుణ్య కేంద్ర నోటీసులు చూడండి.",
            "{need} కోసం పత్రాలు ఇవ్వడానికి లేదా రుసుము చెల్లించడానికి ముందు ప్రభుత్వ బ్యాంకు లేదా నమోదిత కేంద్రంతో నిర్ధారించండి.",
            "మీ ఆదాయాన్ని కాపాడుకోవడానికి పని, చెల్లింపులు, కస్టమర్ ఆర్డర్లను నమోదు చేసుకోండి.",
        ],
    },
    "women": {
        "hi": [
            "{interest} के लिए महिला सहायता सूचना: अगली स्थानीय सेवा तारीख ASHA, ANM, आंगनवाड़ी या पंचायत कार्यकर्ता से पूछें।",
            "{group} जैसे स्वयं सहायता समूह के लिए बैठक, बचत और ऋण का रिकॉर्ड सुरक्षित रखें।",
            "तत्काल खतरे या हिंसा में किसी भरोसेमंद व्यक्ति और आधिकारिक सहायता सेवा से संपर्क करें।",
        ],
        "te": [
            "{interest} కోసం మహిళా సహాయ సమాచారం: తదుపరి స్థానిక సేవ తేదీని ఆశా, ఏఎన్ఎం, అంగన్‌వాడీ లేదా పంచాయతీ సిబ్బందిని అడగండి.",
            "{group} వంటి స్వయం సహాయక సంఘం సమావేశాలు, పొదుపు, రుణ రికార్డులను సురక్షితంగా ఉంచండి.",
            "తక్షణ ప్రమాదం లేదా హింసలో నమ్మకమైన వ్యక్తిని, అధికారిక సహాయ సేవను సంప్రదించండి.",
        ],
    },
    "health": {
        "hi": [
            "{topic}, खासकर {age_group} के लिए स्वास्थ्य सूचना: तारीख सरकारी स्वास्थ्य केंद्र या ASHA कार्यकर्ता से पक्की करें।",
            "सांस लेने में गंभीर परेशानी, सीने में दर्द, अधिक रक्तस्राव, बेहोशी या गंभीर चोट में इलाज में देरी न करें।",
            "दवाइयों और टीकाकरण के रिकॉर्ड साथ रखें और निजी स्वास्थ्य जानकारी सार्वजनिक न करें।",
        ],
        "te": [
            "{topic}, ముఖ్యంగా {age_group} కోసం ఆరోగ్య సమాచారం: తేదీలను ప్రభుత్వ ఆరోగ్య కేంద్రం లేదా ఆశా సిబ్బందితో నిర్ధారించండి.",
            "తీవ్రమైన శ్వాస ఇబ్బంది, ఛాతీ నొప్పి, అధిక రక్తస్రావం, స్పృహ కోల్పోవడం లేదా తీవ్రమైన గాయంలో చికిత్స ఆలస్యం చేయకండి.",
            "మందులు, టీకా రికార్డులను కలిసి ఉంచండి; వ్యక్తిగత ఆరోగ్య వివరాలను బహిరంగంగా పంచుకోకండి.",
        ],
    },
    "education": {
        "hi": [
            "{learner} के लिए शिक्षा सूचना: {need} की जानकारी स्कूल या आधिकारिक पोर्टल पर देखें।",
            "छात्रवृत्ति और अंतिम तारीख के लिए शिक्षक या शिक्षा कार्यालय से पूछें; अप्रमाणित एजेंट को पैसे न दें।",
            "नियमित पढ़ाई का समय तय करें और स्कूल के दस्तावेज व आवेदन नंबर सुरक्षित रखें।",
        ],
        "te": [
            "{learner} కోసం విద్యా సమాచారం: {need} వివరాలను పాఠశాల లేదా అధికారిక పోర్టల్‌లో చూడండి.",
            "స్కాలర్‌షిప్‌లు, చివరి తేదీల కోసం ఉపాధ్యాయుడిని లేదా విద్యా కార్యాలయాన్ని అడగండి; ధృవీకరించని ఏజెంట్‌కు డబ్బు ఇవ్వకండి.",
            "నిరంతర చదువు సమయాన్ని కేటాయించి పాఠశాల పత్రాలు, దరఖాస్తు నంబర్లను భద్రంగా ఉంచండి.",
        ],
    },
    "public_services": {
        "hi": [
            "{ward} में {service} की सार्वजनिक सेवा सूचना: नवीनतम सूचना पंचायत कार्यालय से पक्की करें।",
            "आवेदन की रसीद और पावती नंबर रखें; सरकारी सेवा के लिए अनधिकृत शुल्क न दें।",
            "असुरक्षित सड़क, बिजली या पानी की समस्या आधिकारिक शिकायत माध्यम से बताएं।",
        ],
        "te": [
            "{ward}లో {service} ప్రజా సేవ సమాచారం: తాజా ప్రకటనను పంచాయతీ కార్యాలయంలో నిర్ధారించండి.",
            "దరఖాస్తు రసీదు, స్వీకరణ నంబర్లను ఉంచండి; ప్రభుత్వ సేవలకు అనధికార రుసుము చెల్లించకండి.",
            "ప్రమాదకర రహదారులు, విద్యుత్ లేదా నీటి సమస్యలను అధికారిక ఫిర్యాదు మార్గంలో తెలియజేయండి.",
        ],
    },
    "village_updates": {
        "hi": [
            "{interests} के लिए गांव की सूचना: प्रमाणित पंचायत, स्कूल, स्वास्थ्य केंद्र और विभागीय नोटिस देखें।",
            "संदेश आगे भेजने से पहले तारीख और स्रोत जांचें; अफवाह या अप्रमाणित आपात सूचना न भेजें।",
            "जरूरी स्थानीय हेल्पलाइन नंबर सुरक्षित रखें और केवल प्रमाणित जरूरी सूचना साझा करें।",
        ],
        "te": [
            "{interests} కోసం గ్రామ సమాచారం: ధృవీకరించిన పంచాయతీ, పాఠశాల, ఆరోగ్య కేంద్రం, శాఖల ప్రకటనలను చూడండి.",
            "సందేశాన్ని పంపే ముందు తేదీ, మూలాన్ని తనిఖీ చేయండి; పుకార్లు లేదా ధృవీకరించని అత్యవసర సమాచారాన్ని పంపకండి.",
            "ముఖ్యమైన స్థానిక హెల్ప్‌లైన్ నంబర్లను భద్రపరచి ధృవీకరించిన అత్యవసర సమాచారాన్ని మాత్రమే పంచుకోండి.",
        ],
    },
}


HUB_DATA_DIR = Path(__file__).parent / "data"


def _paths(folder: Path) -> tuple[Path, Path]:
    data_dir = HUB_DATA_DIR
    data_dir.mkdir(exist_ok=True)
    workbook = data_dir / f"{folder.name}.xlsx"
    legacy_workbook = folder / "data" / workbook.name
    if not workbook.exists() and legacy_workbook.exists():
        shutil.copy2(legacy_workbook, workbook)
        print(f"[domain_runtime] Migrated {legacy_workbook} to {workbook}.")
    return workbook, data_dir / f"{folder.name}_outbox.log"


PROFILE_HEADERS = [
    "ID", "Name", "Location", "Phone", "Channel", "Language", "LastSent",
    "CattleType", "Animals", "Count", "Concern", "Source", "Crop", "IrrigationMethod", "Equipment", "Stage",
    "WorkType", "Work", "Need", "Interest", "SupportType", "Group", "Topic", "AgeGroup", "SchoolName", "Learner", "AiFocus", "Service",
    "Ward", "Interests",
]
DELIVERY_HEADERS = ["Timestamp", "ProfileID", "Name", "Location", "Phone", "Channel", "Language", "Message", "Status"]


def _normalized_key(header: str) -> str:
    alias_map = {
        "ID": "id",
        "Name": "name",
        "Location": "location",
        "Phone": "phone",
        "Channel": "channel",
        "Language": "language",
        "LastSent": "last_sent",
        "CattleType": "cattle_type",
        "Animals": "animals",
        "Count": "count",
        "Concern": "concern",
        "Source": "source",
        "Crop": "crop",
        "IrrigationMethod": "irrigation_method",
        "Equipment": "equipment",
        "Stage": "stage",
        "WorkType": "work_type",
        "Work": "work",
        "Need": "need",
        "Interest": "interest",
        "SupportType": "support_type",
        "Group": "group",
        "Topic": "topic",
        "AgeGroup": "age_group",
        "SchoolName": "school_name",
        "Learner": "learner",
        "AiFocus": "ai_focus",
        "Service": "service",
        "Ward": "ward",
        "Interests": "interests",
    }
    return alias_map.get(header, header.lower())


def _profile_values(profile: dict) -> list:
    values = []
    for header in PROFILE_HEADERS:
        values.append(profile.get(_normalized_key(header)))
    return values


def _ensure_workbook(path: Path) -> None:
    if path.exists():
        return
    workbook = Workbook()
    profiles = workbook.active
    profiles.title = "Profiles"
    profiles.append(PROFILE_HEADERS)
    deliveries = workbook.create_sheet("DeliveryLog")
    deliveries.append(DELIVERY_HEADERS)
    legacy = path.with_name(f"{path.stem}.json")
    if legacy.exists():
        try:
            for item in json.loads(legacy.read_text(encoding="utf-8")):
                profiles.append(_profile_values(item))
        except (OSError, json.JSONDecodeError):
            pass
    workbook.save(path)


def _load(path: Path) -> list[dict]:
    _ensure_workbook(path)
    workbook = load_workbook(path, read_only=True)
    sheet = workbook["Profiles"]
    headers = [cell.value for cell in sheet[1]]
    profiles = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        profile = {}
        for header, value in zip(headers, row):
            profile[_normalized_key(header)] = value
        profiles.append(profile)
    workbook.close()
    return profiles


def _save_profiles(path: Path, profiles: list[dict]) -> None:
    workbook = load_workbook(path)
    sheet = workbook["Profiles"]
    sheet.delete_rows(2, sheet.max_row)
    for profile in profiles:
        sheet.append(_profile_values(profile))
    workbook.save(path)


def _log_delivery(path: Path, profile: dict, message: str, timestamp: str) -> None:
    workbook = load_workbook(path)
    workbook["DeliveryLog"].append([
        timestamp, profile["id"], profile["name"], profile["location"],
        profile["phone"], profile["channel"], profile["language"], message, "dry-run-sent",
    ])
    workbook.save(path)


def register(domain: str, folder: Path) -> None:
    definition = DOMAIN_DEFINITIONS[domain]
    profiles_file, _ = _paths(folder)
    profiles = _load(profiles_file)
    profile = {
        "id": max((item["id"] for item in profiles), default=0) + 1,
        "name": input("Name: ").strip(),
        "location": input("Village or location: ").strip(),
        "phone": input("Phone number: ").strip(),
        "channel": input("Channel (sms/whatsapp): ").strip().lower() or "sms",
        "language": input("Language (english/hindi/telugu): ").strip().lower() or "english",
        "last_sent": None,
    }
    for key, prompt in definition["fields"]:
        profile[key] = input(prompt).strip()
    profiles.append(profile)
    _save_profiles(profiles_file, profiles)
    print(f"Saved {definition['title']} profile #{profile['id']}. Run main.py to send updates.")


def _message(domain: str, profile: dict) -> str:
    definition = DOMAIN_DEFINITIONS[domain]
    language = profile.get("language", "english").lower()
    language = {"hindi": "hi", "telugu": "te", "english": "en"}.get(language, language)
    greeting = {"hi": "नमस्ते", "te": "నమస్కారం"}.get(language, "Hello")
    title = {"hi": "दैनिक अपडेट", "te": "రోజువారీ సమాచారం"}.get(language, f"{definition['title']} daily update")
    templates = definition["messages"] if language == "en" else LOCALIZED_MESSAGES[domain].get(language, definition["messages"])
    if domain == "cattle":
        from cattle_agent.nutrition_guide import get_nutrition_plan
        nutrition_note = get_nutrition_plan(profile.get("cattle_type") or profile.get("animals") or "cow", profile.get("concern") or "general")
        templates = list(templates)
        templates.insert(1, f"Feeding note: {nutrition_note}")
    if domain == "irrigation":
        from irrigation_agent.irrigation_guide import get_irrigation_plan
        irrigation_note = get_irrigation_plan(
            profile.get("source") or "borewell",
            profile.get("crop") or "field crop",
            profile.get("irrigation_method") or profile.get("equipment") or "drip",
            profile.get("stage") or "general",
        )
        templates = list(templates)
        templates.insert(2, f"Field schedule: {irrigation_note}")
    if domain == "education":
        from education_agent.ai_guide import get_ai_update
        ai_note = get_ai_update(
            profile.get("school_name") or "School",
            profile.get("learner") or "student",
            profile.get("ai_focus") or "AI basics",
        )
        templates = list(templates)
        templates.insert(1, f"AI school update: {ai_note}")
    if domain == "livelihoods":
        from livelihoods_agent.livelihoods_guide import get_livelihood_update
        livelihood_note = get_livelihood_update(
            profile.get("work_type") or profile.get("work") or "labour",
            profile.get("need") or "jobs",
            profile.get("interest") or "general",
        )
        templates = list(templates)
        templates.insert(1, f"Local opportunity note: {livelihood_note}")
    if domain == "women":
        from women_support.support_guide import get_support_update
        support_note = get_support_update(
            profile.get("interest") or profile.get("support_type") or "health",
            profile.get("group") or "community group",
            profile.get("support_type") or "safety",
        )
        templates = list(templates)
        templates.insert(1, f"Support note: {support_note}")
    lines = [f"{greeting} {profile['name']} ({profile['location']}), VillageNova {title}:"]
    lines.extend(message.format(**profile) for message in templates)
    return "\n".join(lines)


def message_for(domain: str, resident: dict) -> str:
    """Compose a domain update from a hub resident without storing hub data here."""
    if domain == "farmer":
        return FARMER_UPDATE
    profile = dict(resident)
    profile.setdefault("location", "your village")
    for key, _ in DOMAIN_DEFINITIONS[domain]["fields"]:
        profile.setdefault(key, "your local area")
    return _message(domain, profile)


def _generated_message(domain: str, profile: dict) -> str | None:
    """Generate a fresh message when an OpenAI-compatible API is configured."""
    api_key = os.getenv("OPENAI_API_KEY")
    base_url = os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1")
    local_provider = "localhost" in base_url or "127.0.0.1" in base_url
    if not api_key and not local_provider:
        return None
    language = profile.get("language", "english")
    details = ", ".join(
        f"{key}={value}" for key, value in profile.items()
        if key not in {"id", "phone", "channel", "last_sent"} and value
    )
    prompt = (
        f"Create one concise, practical {domain} village update for this resident: {details}. "
        f"Write entirely in {language}. Use the resident's name and location. "
        "Give 3 short actionable points. Do not invent dates, prices, schemes, phone numbers, "
        "or official claims. Tell the resident to verify local information with the relevant "
        "official worker or office. Do not use markdown headings."
    )
    try:
        response = requests.post(
            base_url + "/chat/completions",
            headers={"Content-Type": "application/json", **({"Authorization": f"Bearer {api_key}"} if api_key else {})},
            json={
                "model": os.getenv("OPENAI_MODEL", "gpt-4o-mini"),
                "temperature": 0.8,
                "messages": [
                    {"role": "system", "content": "You write safe, local, useful village service messages."},
                    {"role": "user", "content": prompt},
                ],
            },
            timeout=20,
        )
        response.raise_for_status()
        content = response.json()["choices"][0]["message"]["content"].strip()
        return content or None
    except (requests.RequestException, KeyError, IndexError, TypeError, ValueError):
        return None


def run(domain: str, folder: Path) -> None:
    definition = DOMAIN_DEFINITIONS[domain]
    profiles_file, outbox_file = _paths(folder)
    profiles = _load(profiles_file)
    if not profiles:
        print("No profiles registered. Run agent.py first.")
        return
    send_every_run = os.getenv("SEND_EVERY_RUN", "false").lower() == "true"
    for profile in profiles:
        if not send_every_run and profile.get("last_sent") and datetime.now() - datetime.fromisoformat(profile["last_sent"]) < timedelta(hours=24):
            print(f"{profile['name']}: skipped (message sent within 24 hours).")
            continue
        message = _generated_message(domain, profile) or _message(domain, profile)
        timestamp = datetime.now().isoformat(timespec="seconds")
        with outbox_file.open("a", encoding="utf-8") as stream:
            stream.write(f"[{timestamp}] {profile['channel'].upper()} -> {profile['phone']}: {message}\n")
        _log_delivery(profiles_file, profile, message, timestamp)
        profile["last_sent"] = timestamp
        print(f"[DRY RUN] {profile['channel'].upper()} to {profile['phone']}:\n{message}\n")
    _save_profiles(profiles_file, profiles)
