"""
Excel-backed storage for registered farmers and the delivery log.

Everything lives in one workbook, data/farmers.xlsx, with two sheets:
  - "Farmers"      one row per registered farmer
  - "DeliveryLog"  one row per message sent (or dry-run logged)

Swap this module out for a real database later (SQLite/Postgres) without
touching agent.py or main.py — they only call the functions below.
"""
import os
import shutil
from datetime import datetime
from pathlib import Path
from zipfile import BadZipFile
from openpyxl import Workbook, load_workbook

DATA_DIR = str(Path(__file__).parents[1] / "village_hub" / "data")
FARMERS_FILE = os.path.join(DATA_DIR, "farmers.xlsx")

FARMER_HEADERS = ["ID", "Name", "Village", "Crop", "Stage", "Language", "Phone", "Channel", "Latitude", "Longitude", "LocationSource"]
LOG_HEADERS = ["Timestamp", "FarmerID", "Farmer", "Village", "Channel", "Message", "Status"]


def _create_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Farmers"
    ws.append(FARMER_HEADERS)
    log_ws = wb.create_sheet("DeliveryLog")
    log_ws.append(LOG_HEADERS)
    wb.save(FARMERS_FILE)


def _ensure_workbook():
    os.makedirs(DATA_DIR, exist_ok=True)
    if not os.path.exists(FARMERS_FILE):
        legacy_file = os.path.join(os.path.dirname(__file__), "data", "farmers.xlsx")
        if os.path.exists(legacy_file):
            shutil.copy2(legacy_file, FARMERS_FILE)
            print(f"[storage] Migrated {legacy_file} to {FARMERS_FILE}.")
            return
        _create_workbook()
        return
    try:
        wb = load_workbook(FARMERS_FILE, read_only=True)
        wb.close()
    except (BadZipFile, OSError):
        backup = os.path.splitext(FARMERS_FILE)[0] + f".corrupt-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
        shutil.move(FARMERS_FILE, backup)
        print(f"[storage] Invalid workbook moved to {os.path.basename(backup)}; rebuilding it.")
        _create_workbook()


def _ensure_headers(ws, headers: list[str]):
    existing = [cell.value for cell in ws[1]]
    for header in headers:
        if header not in existing:
            ws.cell(row=1, column=ws.max_column + 1, value=header)
            existing.append(header)


def add_farmer(record: dict) -> int:
    """record keys: name, village, crop, stage, language, phone, channel.
    Returns the new farmer's ID."""
    _ensure_workbook()
    wb = load_workbook(FARMERS_FILE)
    ws = wb["Farmers"]
    _ensure_headers(ws, FARMER_HEADERS)
    next_id = ws.max_row  # header is row 1, so row count so far = next id
    values = {
        "ID": next_id, "Name": record["name"], "Village": record["village"],
        "Crop": record["crop"], "Stage": record["stage"],
        "Language": record["language"], "Phone": record["phone"],
        "Channel": record["channel"], "Latitude": record.get("latitude"),
        "Longitude": record.get("longitude"),
        "LocationSource": record.get("location_source", "manual"),
    }
    ws.append([values.get(header) for header in [cell.value for cell in ws[1]]])
    wb.save(FARMERS_FILE)
    return next_id


def get_all_farmers() -> list[dict]:
    _ensure_workbook()
    wb = load_workbook(FARMERS_FILE)
    ws = wb["Farmers"]
    headers = [cell.value for cell in ws[1]]
    farmers = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        record = dict(zip(headers, row))
        farmers.append({key.lower(): value for key, value in record.items()})
    return farmers


def log_delivery(farmer_name: str, village: str, channel: str, message: str, status: str, farmer_id: int | None = None):
    _ensure_workbook()
    wb = load_workbook(FARMERS_FILE)
    ws = wb["DeliveryLog"]
    _ensure_headers(ws, LOG_HEADERS)
    values = {
        "Timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "FarmerID": farmer_id, "Farmer": farmer_name, "Village": village,
        "Channel": channel, "Message": message, "Status": status,
    }
    ws.append([values.get(header) for header in [cell.value for cell in ws[1]]])
    wb.save(FARMERS_FILE)


def last_delivery_at(farmer_id: int, farmer_name: str | None = None) -> datetime | None:
    """Return the most recent delivery time for a farmer, if one exists."""
    _ensure_workbook()
    wb = load_workbook(FARMERS_FILE, read_only=True)
    ws = wb["DeliveryLog"]
    headers = [cell.value for cell in ws[1]]
    timestamp_index = headers.index("Timestamp")
    farmer_id_index = headers.index("FarmerID") if "FarmerID" in headers else None
    farmer_name_index = headers.index("Farmer") if "Farmer" in headers else None
    latest = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if farmer_id_index is not None:
            if row[farmer_id_index] != farmer_id:
                if row[farmer_id_index] is not None and farmer_name_index is not None:
                    continue
                if farmer_name_index is None or row[farmer_name_index] != farmer_name:
                    continue
        elif farmer_name_index is not None and row[farmer_name_index] != farmer_name:
            continue
        try:
            timestamp = datetime.strptime(str(row[timestamp_index]), "%Y-%m-%d %H:%M:%S")
        except (TypeError, ValueError):
            continue
        if latest is None or timestamp > latest:
            latest = timestamp
    return latest
